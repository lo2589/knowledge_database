"""Turn any supported input into a list of documents, each a list of messages.

A message is {"speaker": "user" | "assistant" | None, "text": Markdown}.
Everything downstream (splitting, display) only ever sees Markdown, so each
format's job is to get its content into Markdown without losing formulas,
code or tables.
"""

from __future__ import annotations

import io
import json
import re
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path


@dataclass
class Document:
    title: str
    format: str
    messages: list[dict] = field(default_factory=list)


class ImportErrorKM(ValueError):
    """Input that cannot be read; the message is shown to the user as-is."""


def load_bytes(name: str, data: bytes) -> list[Document]:
    ext = Path(name).suffix.lower()
    stem = Path(name).stem or "未命名"
    if ext in (".md", ".markdown", ".txt", ".tex", ""):
        return [Document(stem, "markdown", [{"speaker": None, "text": _decode(data)}])]
    if ext in (".html", ".htm"):
        html = _decode(data)
        return [Document(_html_title(html) or stem, "html", [{"speaker": None, "text": html_to_markdown(html)}])]
    if ext == ".jsonl":
        return [claude_code_transcript(_decode(data), stem)]
    if ext == ".json":
        return chat_export(json.loads(_decode(data)), stem)
    if ext == ".pdf":
        return [Document(stem, "pdf", [{"speaker": None, "text": pdf_text(data)}])]
    if ext == ".docx":
        return [Document(stem, "docx", [{"speaker": None, "text": docx_markdown(data)}])]
    if ext == ".ipynb":
        return [Document(stem, "notebook", [{"speaker": None, "text": notebook_markdown(_decode(data))}])]
    if ext == ".rst":
        return [Document(stem, "rst", [{"speaker": None, "text": _decode(data)}])]
    raise ImportErrorKM(f"不支持 {ext} 文件；支持 .md .txt .tex .rst .html .pdf .docx .ipynb .json（ChatGPT/Claude 导出）.jsonl（Claude Code 会话）")


# --- a whole folder (say, a code repository's docs) ---------------------------

FOLDER_EXTS = {".md", ".markdown", ".txt", ".tex", ".rst", ".html", ".htm", ".pdf", ".docx", ".ipynb"}
SKIP_DIRS = {".git", ".hg", ".svn", "node_modules", ".venv", "venv", "__pycache__", "dist", "build",
             "target", ".next", ".cache", ".idea", ".vscode", ".pytest_cache", ".mypy_cache", "site-packages"}


def load_folder(root: str, limit: int = 300, max_bytes: int = 20 * 1024 * 1024) -> tuple[list[Document], list[str]]:
    """Every supported document under `root`, titled by its path inside it.

    Returns (documents, skipped) where skipped says why each file was left out,
    so a partial import is never silent.
    """
    base = Path(root).expanduser()
    if not base.is_dir():
        raise ImportErrorKM(f"{base} 不是文件夹")
    files = []
    for path in sorted(base.rglob("*")):
        rel = path.relative_to(base)
        if any(part in SKIP_DIRS or part.startswith(".") for part in rel.parts[:-1]):
            continue
        if path.is_file() and path.suffix.lower() in FOLDER_EXTS:
            files.append(path)
    if not files:
        raise ImportErrorKM(f"{base} 里没有能导入的文档（{' '.join(sorted(FOLDER_EXTS))}）")
    docs, skipped = [], []
    for path in files[:limit]:
        rel = str(path.relative_to(base))
        if path.stat().st_size > max_bytes:
            skipped.append(f"{rel}：超过 {max_bytes // 1024 // 1024}MB")
            continue
        try:
            for d in load_bytes(path.name, path.read_bytes()):
                d.title = f"{base.name}/{rel}"
                docs.append(d)
        except (ImportErrorKM, ValueError, OSError) as exc:
            skipped.append(f"{rel}：{exc}")
    if len(files) > limit:
        skipped.append(f"还有 {len(files) - limit} 个文件超出单次上限 {limit}，没导")
    return docs, skipped


# --- Jupyter notebooks: markdown cells as-is, code cells fenced ---------------


def notebook_markdown(text: str) -> str:
    nb = json.loads(text)
    lang = ((nb.get("metadata") or {}).get("language_info") or {}).get("name") or "python"
    out = []
    for cell in nb.get("cells", []):
        src = cell.get("source", "")
        src = "".join(src) if isinstance(src, list) else str(src)
        if not src.strip():
            continue
        if cell.get("cell_type") == "markdown":
            out.append(src.strip())
        elif cell.get("cell_type") == "code":
            out.append(f"```{lang}\n{src.rstrip()}\n```")
    return "\n\n".join(out)


def load_text(text: str, title: str = "") -> Document:
    """Pasted text. HTML pasted from a web page is converted; anything else is Markdown."""
    stripped = text.lstrip()
    if stripped.startswith("<") and re.search(r"</(p|div|li|h\d|table|span)>", stripped, re.I):
        return Document(title or _html_title(text) or "粘贴的网页", "html", [{"speaker": None, "text": html_to_markdown(text)}])
    first = next((l.strip("# ").strip() for l in text.splitlines() if l.strip()), "")
    return Document(title or first[:40] or "粘贴的文字", "markdown", [{"speaker": None, "text": text}])


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "gb18030"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


# --- Claude Code session transcripts (~/.claude/projects/*/*.jsonl) ---------

# Harness-injected wrappers inside user messages; what the person typed is
# whatever is left after removing them.
_INJECTED = re.compile(
    r"<(system-reminder|ide_[a-z_]+|command-[a-z-]+|local-command-[a-z-]+|user-prompt-submit-hook)\b[^>]*>.*?</\1>",
    re.S,
)


def claude_code_transcript(text: str, fallback_title: str) -> Document:
    title, messages = "", []
    for line in text.splitlines():
        try:
            rec = json.loads(line)
        except json.JSONDecodeError:
            continue
        if rec.get("type") == "ai-title" and rec.get("aiTitle"):
            title = rec["aiTitle"]
        if rec.get("type") not in ("user", "assistant") or rec.get("isSidechain") or rec.get("isMeta"):
            continue
        msg = rec.get("message") or {}
        content = msg.get("content")
        if isinstance(content, str):
            blocks = [content]
        elif isinstance(content, list):
            # Only what was said: no thinking, tool calls or tool output.
            blocks = [b.get("text", "") for b in content if isinstance(b, dict) and b.get("type") == "text"]
        else:
            continue
        body = _INJECTED.sub("", "\n\n".join(blocks)).strip()
        if not body:
            continue
        speaker = msg.get("role") or rec["type"]
        if messages and messages[-1]["speaker"] == speaker:
            messages[-1]["text"] += "\n\n" + body
        else:
            messages.append({"speaker": speaker, "text": body})
    if not messages:
        raise ImportErrorKM("这个 .jsonl 里没有找到对话内容（不是 Claude Code 会话记录？）")
    return Document(title or fallback_title, "claude-code", messages)


def list_claude_sessions(root: Path | None = None, limit: int = 40) -> list[dict]:
    root = root or Path.home() / ".claude" / "projects"
    files = sorted(root.glob("*/*.jsonl"), key=lambda p: p.stat().st_mtime, reverse=True)[:limit]
    out = []
    for path in files:
        title, first = "", ""
        try:
            with path.open(encoding="utf-8", errors="replace") as fh:
                for line in fh:
                    if '"ai-title"' in line and not title:
                        title = json.loads(line).get("aiTitle", "")
                    elif not first and '"type":"user"' in line and '"tool_result"' not in line:
                        rec = json.loads(line)
                        c = (rec.get("message") or {}).get("content")
                        t = c if isinstance(c, str) else " ".join(
                            b.get("text", "") for b in c or [] if isinstance(b, dict) and b.get("type") == "text")
                        first = _INJECTED.sub("", t).strip()[:80]
                    if title and first:
                        break
        except (OSError, json.JSONDecodeError):
            continue
        out.append({
            "path": str(path),
            "project": path.parent.name.replace("-Users-", "~/").replace("-", "/")[-60:],
            "title": title or first or path.stem,
            "first": first,
            "mtime": int(path.stat().st_mtime * 1000),
        })
    return out


# --- ChatGPT / Claude.ai data exports --------------------------------------


def chat_export(data, fallback_title: str) -> list[Document]:
    convs = data if isinstance(data, list) else [data]
    docs = []
    for conv in convs:
        if not isinstance(conv, dict):
            continue
        if "mapping" in conv:  # ChatGPT conversations.json
            docs.append(_chatgpt(conv, fallback_title))
        elif "chat_messages" in conv:  # Claude.ai export
            msgs = []
            for m in conv["chat_messages"]:
                text = m.get("text") or "\n\n".join(
                    c.get("text", "") for c in m.get("content", []) if c.get("type") == "text")
                if text.strip():
                    msgs.append({"speaker": "user" if m.get("sender") == "human" else "assistant", "text": text})
            docs.append(Document(conv.get("name") or fallback_title, "claude-export", msgs))
    docs = [d for d in docs if d.messages]
    if not docs:
        raise ImportErrorKM("这个 .json 不是 ChatGPT 或 Claude 的对话导出")
    return docs


def _chatgpt(conv: dict, fallback_title: str) -> Document:
    mapping = conv["mapping"]
    # Follow the branch that ends at current_node, so edited/regenerated
    # alternatives do not all show up.
    node, chain = conv.get("current_node"), []
    while node and node in mapping:
        chain.append(mapping[node])
        node = mapping[node].get("parent")
    msgs = []
    for item in reversed(chain):
        m = item.get("message") or {}
        role = (m.get("author") or {}).get("role")
        parts = (m.get("content") or {}).get("parts") or []
        text = "\n\n".join(p for p in parts if isinstance(p, str)).strip()
        if role in ("user", "assistant") and text:
            msgs.append({"speaker": role, "text": _chatgpt_math(text)})
    return Document(conv.get("title") or fallback_title, "chatgpt", msgs)


def _chatgpt_math(text: str) -> str:
    # ChatGPT writes \( \) and \[ \]; the renderer handles both, nothing to do,
    # but some exports double-escape them.
    return text.replace("\\\\(", "\\(").replace("\\\\)", "\\)").replace("\\\\[", "\\[").replace("\\\\]", "\\]")


# --- PDF / Word ---------------------------------------------------------------


def pdf_text(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ImportErrorKM("读 PDF 需要先安装：pip install pypdf") from exc
    reader = PdfReader(io.BytesIO(data))
    pages = [(p.extract_text() or "").strip() for p in reader.pages]
    text = "\n\n".join(p for p in pages if p)
    if not text.strip():
        raise ImportErrorKM("这个 PDF 里提取不到文字（可能是扫描件，需要先 OCR）")
    # PDF lines are hard-wrapped; rejoin lines inside a paragraph.
    return re.sub(r"(?<![。！？.!?:：])\n(?!\n)", " ", text)


def docx_markdown(data: bytes) -> str:
    try:
        import docx
    except ImportError as exc:
        raise ImportErrorKM("读 Word 需要先安装：pip install python-docx") from exc
    document = docx.Document(io.BytesIO(data))
    out = []
    for block in document.element.body.iterchildren():
        tag = block.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = docx.text.paragraph.Paragraph(block, document)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name or "").lower() if para.style is not None else ""
            m = re.match(r"heading (\d)", style)
            if m:
                out.append("#" * int(m.group(1)) + " " + text)
            elif "list" in style:
                out.append("- " + text)
            else:
                out.append(text)
        elif tag == "tbl":
            table = docx.table.Table(block, document)
            rows = [[c.text.strip().replace("|", "\\|").replace("\n", " ") for c in r.cells] for r in table.rows]
            if rows:
                out.append(_md_table(rows))
    return "\n\n".join(out)


def _md_table(rows: list[list[str]]) -> str:
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
    lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(lines)


# --- HTML → Markdown -----------------------------------------------------------


def _html_title(html: str) -> str:
    m = re.search(r"<title[^>]*>(.*?)</title>", html, re.I | re.S)
    return re.sub(r"\s+", " ", m.group(1)).strip() if m else ""


class _H2M(HTMLParser):
    """Small HTML → Markdown converter that keeps math, code and tables.

    Math is recovered from what renderers leave behind: KaTeX's
    <annotation encoding="application/x-tex">, MathJax's
    <script type="math/tex">, and <math alttext>.
    """

    BLOCK = {"p", "div", "section", "article", "header", "footer", "main", "figure", "br", "hr"}
    SKIP = {"script", "style", "noscript", "head", "svg", "button", "nav"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.out: list[str] = []
        self.skip = 0
        self.pre = 0
        self.lists: list[list] = []  # [kind, counter]
        self.href: list[str | None] = []
        self.table: list[list[str]] | None = None
        self.cell: list[str] | None = None
        self.math: dict | None = None  # collecting TeX
        self.math_depth = 0
        self.fence_at = -1

    def w(self, s: str) -> None:
        if self.cell is not None:
            self.cell.append(s)
        else:
            self.out.append(s)

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        cls = a.get("class") or ""
        if self.math is not None:
            self.math_depth += 1
            if tag == "annotation" and "tex" in (a.get("encoding") or ""):
                self.math["collect"] = True
            return
        if tag == "span" and ("katex-display" in cls or cls.split() == ["katex"] or "katex" in cls.split()):
            self.math = {"display": "katex-display" in cls, "tex": [], "collect": False}
            self.math_depth = 1
            return
        if tag == "math":
            tex = a.get("alttext")
            if tex:
                display = a.get("display") == "block"
                self.w(f"\n\n$${tex}$$\n\n" if display else f"${tex}$")
                self.skip += 1
                return
        if tag == "script" and (a.get("type") or "").startswith("math/tex"):
            self.math = {"display": "mode=display" in a.get("type", ""), "tex": [], "collect": True, "script": True}
            self.math_depth = 1
            return
        if tag in self.SKIP:
            self.skip += 1
            return
        if self.skip:
            return
        if tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
            self.w("\n\n" + "#" * int(tag[1]) + " ")
        elif tag in self.BLOCK:
            self.w("\n\n" if tag != "br" else "\n")
        elif tag == "pre":
            self.pre += 1
            lang = ""
            m = re.search(r"language-([\w+-]+)", cls)
            if m:
                lang = m.group(1)
            self.fence_at = len(self.out)
            self.w(f"\n\n```{lang}\n")
        elif tag == "code" and self.pre:
            # Highlighters put the language on <code>, not <pre>.
            m = re.search(r"language-([\w+-]+)", cls)
            if m and self.cell is None and self.out[self.fence_at:self.fence_at + 1] == ["\n\n```\n"]:
                self.out[self.fence_at] = f"\n\n```{m.group(1)}\n"
        elif tag == "code" and not self.pre:
            self.w("`")
        elif tag in ("strong", "b"):
            self.w("**")
        elif tag in ("em", "i"):
            self.w("*")
        elif tag == "a":
            self.href.append(a.get("href"))
            self.w("[")
        elif tag == "img":
            self.w(f"![{a.get('alt', '')}]({a.get('src', '')})")
        elif tag in ("ul", "ol"):
            self.lists.append([tag, 0])
        elif tag == "li":
            depth = max(0, len(self.lists) - 1)
            kind = self.lists[-1] if self.lists else ["ul", 0]
            kind[1] += 1
            marker = f"{kind[1]}." if kind[0] == "ol" else "-"
            self.w("\n" + "  " * depth + marker + " ")
        elif tag == "blockquote":
            self.w("\n\n> ")
        elif tag == "table":
            self.table = []
        elif tag == "tr" and self.table is not None:
            self.table.append([])
        elif tag in ("td", "th") and self.table is not None:
            self.cell = []

    def handle_endtag(self, tag):
        if self.math is not None:
            self.math_depth -= 1
            if tag == "annotation":
                self.math["collect"] = False
            if self.math_depth <= 0:
                tex = "".join(self.math["tex"]).strip()
                if tex:
                    self.w(f"\n\n$$\n{tex}\n$$\n\n" if self.math["display"] else f"${tex}$")
                self.math = None
            return
        if tag in self.SKIP or (tag == "math" and self.skip):
            self.skip = max(0, self.skip - 1)
            return
        if self.skip:
            return
        if tag in ("h1", "h2", "h3", "h4", "h5", "h6", "p", "div", "section", "article", "blockquote"):
            self.w("\n\n")
        elif tag == "pre":
            self.pre = max(0, self.pre - 1)
            self.w("\n```\n\n")
        elif tag == "code" and not self.pre:
            self.w("`")
        elif tag in ("strong", "b"):
            self.w("**")
        elif tag in ("em", "i"):
            self.w("*")
        elif tag == "a":
            href = self.href.pop() if self.href else None
            self.w(f"]({href})" if href else "]")
        elif tag in ("ul", "ol"):
            if self.lists:
                self.lists.pop()
            self.w("\n\n")
        elif tag in ("td", "th") and self.table is not None and self.cell is not None:
            text = re.sub(r"\s+", " ", "".join(self.cell)).strip().replace("|", "\\|")
            self.cell = None
            if self.table:
                self.table[-1].append(text)
        elif tag == "table" and self.table is not None:
            rows = [r for r in self.table if r]
            self.table = None
            if rows:
                self.out.append("\n\n" + _md_table(rows) + "\n\n")

    def handle_data(self, data):
        if self.math is not None:
            if self.math.get("collect"):
                self.math["tex"].append(data)
            return
        if self.skip:
            return
        if self.pre:
            self.w(data)
        else:
            self.w(re.sub(r"\s+", " ", data))


def html_to_markdown(html: str) -> str:
    parser = _H2M()
    parser.feed(html)
    parser.close()
    text = "".join(parser.out)
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
